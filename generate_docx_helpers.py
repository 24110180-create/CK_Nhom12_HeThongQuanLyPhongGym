import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc, title_text):
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"FitLife Gym Pro | {title_text}")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 120, 120)

        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Đồ Án Cuối Kỳ - Nhóm 12 | Hệ Thống Quản Trị & Vận Hành Phòng Gym FitLife")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(120, 120, 120)

def format_doc_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.font.name = "Segoe UI"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Dark Navy

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(18)
    run2 = p2.add_run(subtitle)
    run2.font.name = "Segoe UI"
    run2.font.size = Pt(12)
    run2.font.italic = True
    run2.font.color.rgb = RGBColor(75, 85, 99)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Segoe UI"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138) # Navy
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Segoe UI"
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(13, 148, 136) # Teal
    return h

def add_paragraph(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Segoe UI"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(17, 24, 39)
    r = p.add_run(text)
    r.font.name = "Segoe UI"
    r.font.size = Pt(10.5)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(31, 41, 55)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Segoe UI"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(17, 24, 39)
    r = p.add_run(text)
    r.font.name = "Segoe UI"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(31, 41, 55)
    return p

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(17, 24, 39)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_callout(doc, text, title="LƯU Ý QUAN TRỌNG", border_color="EAB308", bg_color="FEFCE8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(146, 64, 14)
    
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(113, 63, 18)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def create_table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], "1E3A8A") # Navy
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Segoe UI"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data rows
    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=130, right=130)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = "Segoe UI"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(31, 41, 55)
                
    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

print("Helper functions defined successfully")
